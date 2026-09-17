"""从简历结构化 JSON 直接生成 .docx 文档（不依赖 AI、不依赖原始 docx 模板）。

用于「导出 Word 简历」功能：把优化后的 structured JSON 渲染为一份
单栏、ATS 友好的 Word 文档，便于用户在 Word 中二次编辑。
"""
from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def build_docx_from_resume_json(data: dict) -> bytes:
    """根据简历 JSON 生成 docx 二进制内容。

    data 结构（与 parsed_json / optimized_json 一致）：
      personal_info: {name, email, phone, ...}
      summary: str
      experience: [{title, company, start, end, points: [...]}]
      education: [{school, degree, major, start, end}]
      skills: [str]
      projects: [{name, description, tech: [...]}]
    """
    data = data or {}
    doc = Document()

    # 基础样式：字体、字号
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    def _add_heading(text: str):
        h = doc.add_heading(text, level=1)
        for run in h.runs:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        return h

    # 头部：姓名 + 联系方式
    info = data.get("personal_info") or {}
    name = info.get("name") or "简历"
    title = doc.add_heading(name, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    contact_parts = []
    if info.get("phone"):
        contact_parts.append(f"电话：{info['phone']}")
    if info.get("email"):
        contact_parts.append(f"邮箱：{info['email']}")
    for extra in ("location", "city", "current_title"):
        if info.get(extra):
            contact_parts.append(str(info[extra]))
    if contact_parts:
        doc.add_paragraph("  |  ".join(contact_parts))

    # 个人总结
    summary = data.get("summary")
    if summary:
        _add_heading("个人总结")
        doc.add_paragraph(summary)

    # 工作经历
    experiences = data.get("experience") or []
    if experiences:
        _add_heading("工作经历")
        for exp in experiences:
            exp = exp or {}
            line = f"{exp.get('title', '')}  @  {exp.get('company', '')}"
            period = f"{exp.get('start', '')} - {exp.get('end', '')}"
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            if period:
                p.add_run(f"    （{period}）")
            for point in exp.get("points") or []:
                if point:
                    doc.add_paragraph(str(point), style="List Bullet")

    # 项目经历
    projects = data.get("projects") or []
    if projects:
        _add_heading("项目经历")
        for proj in projects:
            proj = proj or {}
            p = doc.add_paragraph()
            run = p.add_run(proj.get("name", ""))
            run.bold = True
            if proj.get("description"):
                doc.add_paragraph(proj["description"])
            tech = proj.get("tech") or []
            if tech:
                doc.add_paragraph("技术栈：" + ", ".join(str(t) for t in tech))

    # 教育背景
    educations = data.get("education") or []
    if educations:
        _add_heading("教育背景")
        for edu in educations:
            edu = edu or {}
            parts = [edu.get("school", ""), edu.get("degree", ""), edu.get("major", "")]
            line = "  |  ".join(str(p) for p in parts if p)
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            period = f"{edu.get('start', '')} - {edu.get('end', '')}"
            if period.strip(" -"):
                p.add_run(f"    （{period}）")

    # 技能
    skills = data.get("skills") or []
    if skills:
        _add_heading("专业技能")
        doc.add_paragraph("、".join(str(s) for s in skills))

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
